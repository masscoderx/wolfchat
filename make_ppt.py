from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(18, 53, 48)
GREEN = RGBColor(56, 133, 110)
TEAL = RGBColor(98, 176, 161)
TEXT = RGBColor(34, 39, 35)
MUTED = RGBColor(89, 102, 96)
LIGHT = RGBColor(248, 247, 242)
WHITE = RGBColor(255, 255, 255)
ACCENT = RGBColor(230, 126, 88)
SOFT = RGBColor(230, 239, 234)
DARK = RGBColor(16, 25, 23)


def solid_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_header_line(slide, y=0.55, length=2.0):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(y), Inches(length), Inches(0.08))
    line.fill.solid(); line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def add_title_slide(title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(slide, LIGHT)
    # subtle decorative blocks
    for x, y, w, h, c in [
        (0.7, 0.6, 1.4, 1.1, SOFT),
        (10.8, 0.9, 1.5, 1.0, SOFT),
        (11.4, 5.8, 1.1, 0.9, SOFT),
    ]:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = c
        shape.line.fill.background()

    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.9), Inches(2.0), Inches(0.55))
    badge.fill.solid(); badge.fill.fore_color.rgb = NAVY
    badge.line.fill.background()
    badge_tf = badge.text_frame; badge_tf.text = 'WOLFCHAT'
    badge_tf.paragraphs[0].alignment = 1
    badge_tf.paragraphs[0].font.bold = True
    badge_tf.paragraphs[0].font.size = Pt(13)
    badge_tf.paragraphs[0].font.color.rgb = WHITE

    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(8.8), Inches(1.4))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY

    subtitle_box = slide.shapes.add_textbox(Inches(0.95), Inches(3.2), Inches(8.7), Inches(1.1))
    tf2 = subtitle_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(17)
    p2.font.color.rgb = MUTED

    # floating stat cards
    for i, (label, value, x, y, c) in enumerate([
        ('Goals', '17', 8.8, 2.0, GREEN),
        ('Impact', 'SDGs', 9.9, 3.2, NAVY),
        ('Privacy', 'Local', 8.8, 4.4, ACCENT),
    ]):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.3), Inches(1.0))
        box.fill.solid(); box.fill.fore_color.rgb = c
        box.line.fill.background()
        tfb = box.text_frame; tfb.word_wrap = True; tfb.paragraphs[0].text = label
        tfb.paragraphs[0].font.size = Pt(10); tfb.paragraphs[0].font.bold = True; tfb.paragraphs[0].font.color.rgb = WHITE
        pval = tfb.add_paragraph(); pval.text = value; pval.font.size = Pt(21); pval.font.bold = True; pval.font.color.rgb = WHITE; pval.alignment = 1


def add_slide_header(slide, title):
    solid_bg(slide, LIGHT)
    add_header_line(slide, 0.5, 2.2)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(9.5), Inches(0.7))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = NAVY


def add_bullets_slide(title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, title)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.5), Inches(4.8))
    box.fill.solid(); box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = RGBColor(219, 228, 222)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.35); tf.margin_right = Inches(0.3)
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT
        p.space_after = Pt(8)


def add_feature_cards(title, cards):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, title)
    positions = [
        (0.9, 1.8, 5.1, 2.2),
        (6.3, 1.8, 5.6, 2.2),
        (0.9, 4.2, 5.1, 2.2),
        (6.3, 4.2, 5.6, 2.2),
    ]
    for idx, (label, text) in enumerate(cards):
        x, y, w, h = positions[idx]
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        card.fill.solid(); card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(219, 228, 222)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.2); tf.margin_top = Inches(0.12)
        title_p = tf.paragraphs[0]
        title_p.text = label
        title_p.font.size = Pt(17)
        title_p.font.bold = True
        title_p.font.color.rgb = NAVY
        body_p = tf.add_paragraph()
        body_p.text = text
        body_p.font.size = Pt(12)
        body_p.font.color.rgb = MUTED


def add_process_slide(title, steps):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, title)
    for i, step in enumerate(steps):
        x = 0.8 + i * 2.55
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(2.1), Inches(1.1), Inches(1.1))
        circle.fill.solid(); circle.fill.fore_color.rgb = GREEN if i % 2 == 0 else NAVY
        circle.line.fill.background()
        tf = circle.text_frame; tf.text = str(i+1); tf.paragraphs[0].alignment = 1; tf.paragraphs[0].font.size = Pt(20); tf.paragraphs[0].font.bold = True; tf.paragraphs[0].font.color.rgb = WHITE
        txt = slide.shapes.add_textbox(Inches(x-0.1), Inches(3.5), Inches(1.5), Inches(1.6))
        t = txt.text_frame; t.word_wrap=True; p=t.paragraphs[0]; p.text=step; p.font.size=Pt(11); p.font.color.rgb=TEXT


def add_sdg_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, 'How Wolfchat Supports the SDGs')
    goals = [
        ('SDG 4', 'Quality Education', '#2E7D66'),
        ('SDG 8', 'Decent Work', '#E27C57'),
        ('SDG 9', 'Innovation', '#4D6CC8'),
        ('SDG 12', 'Responsible Production', '#5FAE9B'),
        ('SDG 13', 'Climate Action', '#2F7C6E'),
        ('SDG 17', 'Partnerships', '#7E9BA7'),
    ]
    positions = [
        (0.8, 1.9, 2.2, 1.3), (3.2, 1.9, 2.2, 1.3), (5.6, 1.9, 2.2, 1.3),
        (8.0, 1.9, 2.2, 1.3), (10.4, 1.9, 2.2, 1.3), (0.8, 3.6, 2.2, 1.3),
        (3.2, 3.6, 2.2, 1.3), (5.6, 3.6, 2.2, 1.3), (8.0, 3.6, 2.2, 1.3),
        (10.4, 3.6, 2.2, 1.3), (0.8, 5.3, 2.2, 1.3), (3.2, 5.3, 2.2, 1.3),
    ]
    for idx, (goal, title, color) in enumerate(goals):
        if idx >= len(positions):
            break
        x, y, w, h = positions[idx]
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor.from_string(color) if False else RGBColor(0,0,0)
        # custom hex conversion
        hex_color = color.lstrip('#')
        rgb = tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
        box.fill.fore_color.rgb = RGBColor(*rgb)
        box.line.fill.background()
        tf = box.text_frame; tf.word_wrap = True; tf.margin_left=Inches(0.12); tf.margin_right=Inches(0.12)
        p1 = tf.paragraphs[0]; p1.text = goal; p1.font.size=Pt(8); p1.font.bold=True; p1.font.color.rgb = WHITE
        p2 = tf.add_paragraph(); p2.text = title; p2.font.size=Pt(8.5); p2.font.bold=True; p2.font.color.rgb = WHITE

    bottom = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.5), Inches(0.5))
    p = bottom.text_frame.paragraphs[0]
    p.text = 'Wolfchat helps teams turn data and documents into actionable progress monitoring for a more sustainable future.'
    p.font.size = Pt(11)
    p.font.color.rgb = MUTED

# Slide 1
add_title_slide(
    'Wolfchat',
    'Private local AI for research, evidence-based decisions, and SDG-aligned action.'
)

# Slide 2
add_bullets_slide(
    'The Challenge',
    [
        'Teams collect information in scattered files, PDFs, notes, and spreadsheets.',
        'Current AI helpers are often generic, disconnected from source material, or not private.',
        'Decision-makers need fast answers, grounded evidence, and clear action paths.',
        'Without this, research becomes slow, inconsistent, and riskier than it should be.'
    ]
)

# Slide 3
add_feature_cards(
    'Why Wolfchat Is Different',
    [
        ('Private by design', 'Files stay local in the browser; nothing is sent to a hosted API by default.'),
        ('Evidence-backed answers', 'The app matches questions to source passages and shows supporting excerpts.'),
        ('Smart data views', 'CSV files can be analyzed and visualized as bar or line charts.'),
        ('SDG dashboard', 'Goal tracking and sustainability-focused insight help connect action to global impact.')
    ]
)

# Slide 4
add_bullets_slide(
    'Core Product Features',
    [
        'General chat mode for everyday explanations, planning, and writing support.',
        'File-aware mode that answers using uploaded evidence and matched passages.',
        'Knowledge map for topic understanding, contradiction checks, and evidence status.',
        'Citations and direct source receipts to build trust and reduce hallucination risk.'
    ]
)

# Slide 5
add_process_slide(
    'How It Works',
    [
        'Upload files',
        'Index content',
        'Match evidence',
        'Answer with citations',
        'Visualize trends',
        'Track impact'
    ]
)

# Slide 6
add_sdg_slide()

# Slide 7
add_bullets_slide(
    'Real-World Use Cases',
    [
        'Researchers and students can summarize documents, ask questions, and verify references quickly.',
        'NGOs and public institutions can track SDG-related priorities and action plans from local evidence.',
        'Teams can analyze CSV datasets, discover trends, and compare claims using visual insights.',
        'Decision-makers can improve clarity, reduce uncertainty, and move from information overload to action.'
    ]
)

# Slide 8
add_bullets_slide(
    'Future Potential',
    [
        'Add stronger multi-document summarization and cross-file comparison.',
        'Introduce more advanced SDG scoring, recommendation engines, and reporting templates.',
        'Expand data forecasting, trend analysis, and personalized dashboards.',
        'Scale for schools, NGOs, startups, and public sector institutions with privacy-first deployment.'
    ]
)

# Closing slide
add_title_slide(
    'Conclusion',
    'Wolfchat combines AI, local privacy, evidence-based research, and SDG-focused decision support in one powerful platform.'
)

ppt_file = r'C:\Users\hp\OneDrive\imb main\Wolfchat_SDG_Presentation.pptx'
prs.save(ppt_file)
print(ppt_file)
